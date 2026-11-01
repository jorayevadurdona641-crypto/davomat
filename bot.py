import asyncio
import logging
import os
import math
from datetime import datetime
from aiogram import Bot, Dispatcher, F
from aiogram.filters import Command
from aiogram.types import Message, FSInputFile, ReplyKeyboardMarkup, KeyboardButton
from aiogram.fsm.state import StatesGroup, State
from aiogram.fsm.context import FSMContext
import openpyxl

# --- SOZLAMALAR ---
API_TOKEN = '8665708174:AAF9WFb0j0...'  # Bot tokeningizni yozing

# 4 ta adminning Telegram ID raqamlari
ADMIN_IDS = [1179846016, 1838498079, 1611805500, 727029645]

# Texnikumning GPS koordinatalari (Yakkabog' uchun)
COLLEGE_LAT = 38.9333
COLLEGE_LON = 66.8167
MAX_DISTANCE_METERS = 300  # Ruxsat etilgan radius (metrda)

bot = Bot(token=API_TOKEN)
dp = Dispatcher()

EXCEL_DB = "Umumiy ma'lumot.xlsx"  # Rasmda ko'rsatilgan Excel fayl nomi
LOG_FILE = "attendance_log.csv"
EXCEL_REPORT = "xodimlar_hisoboti.xlsx"

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")

class AttendanceState(StatesGroup):
    waiting_for_jshr = State()
    waiting_for_location = State()
    waiting_for_photo = State()
    waiting_for_leave = State()

# Excel fayldan PNFL (JSHR) bo'yicha qidirish va farrosh/qorovullarni o'tkazib yuborish
def get_teacher_from_excel(jshr: str):
    if not os.path.exists(EXCEL_DB):
        logging.error(f"{EXCEL_DB} fayli topilmadi!")
        return None
    try:
        wb = openpyxl.load_workbook(EXCEL_DB, data_only=True)
        ws = wb.active  # 1-sahifa
        
        # 4-qatordan boshlab o'qiymiz (chunki 1, 2, 3-qatorlarda sarlavhalar bor)
        for row in ws.iter_rows(min_row=4, values_only=True):
            if row and len(row) >= 5:
                # Rasmga ko'ra ustunlar:
                # row[1] (B ustuni) -> Ism sharifi
                # row[2] (C ustuni) -> Lavozimi
                # row[4] (E ustuni) -> PNFL (JSHR)
                
                row_name = str(row[1]).strip() if row[1] is not None else ""
                row_position = str(row[2]).strip().lower() if row[2] is not None else ""
                row_pnfl = str(row[4]).strip() if row[4] is not None else ""
                
                # Agar JSHR mos kelsa
                if row_pnfl == str(jshr).strip():
                    # Farrosh yoki qorovullarni tekshirib tashlab yuboramiz
                    if "farrosh" in row_position or "qorovul" in row_position:
                        return "EXCLUDED"  # Farrosh yoki qorovul bo'lsa
                    return row_name
    except Exception as e:
        logging.error(f"Excel faylni o'qishda xatolik: {e}")
    return None

def calculate_distance(lat1, lon1, lat2, lon2):
    R = 6371000
    phi1 = math.radians(lat1)
    phi2 = math.radians(lat2)
    delta_phi = math.radians(lat2 - lat1)
    delta_lambda = math.radians(lon2 - lon1)
    a = math.sin(delta_phi / 2)**2 + math.cos(phi1) * math.cos(phi2) * math.sin(delta_lambda / 2)**2
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return R * c

# Ishga kelish vaqtini tekshirish (Chegara: 08:30)
def check_arrival_status():
    now = datetime.now()
    limit_time = now.replace(hour=8, minute=30, second=0, microsecond=0)
    if now > limit_time:
        return "Kechikib keldi"
    else:
        return "O'z vaqtida keldi"

@dp.message(Command("start"))
async def cmd_start(message: Message, state: FSMContext):
    await state.set_state(AttendanceState.waiting_for_jshr)
    await message.answer("Assalomu alaykum! Davomat tizimi. PNFL (JSHR) raqamingizni kiriting:")

@dp.message(AttendanceState.waiting_for_jshr, F.text)
async def process_jshr(message: Message, state: FSMContext):
    jshr = message.text.strip()
    result = get_teacher_from_excel(jshr)  # Excel fayldan tekshirish
    
    if result == "EXCLUDED":
        await message.answer("❌ Kechirasiz, ushbu tizim orqali davomat qayd etish qoidalariga ko'ra sizning lavozimingiz kiritilmagan.")
        await state.clear()
    elif result:
        await state.update_data(jshr=jshr, teacher_name=result)
        kb = ReplyKeyboardMarkup(
            keyboard=[[KeyboardButton(text="📍 Geolokatsiyani yuborish", request_location=True)]],
            resize_keyboard=True,
            one_time_keyboard=True
        )
        await state.set_state(AttendanceState.waiting_for_location)
        await message.answer(f"Salom, **{result}**! Lokatsiyangizni yuboring:", reply_markup=kb, parse_mode="Markdown")
    else:
        await message.answer("Bunday PNFL (JSHR) 'Umumiy ma'lumot.xlsx' faylidan topilmadi. Qaytadan kiriting:")

@dp.message(AttendanceState.waiting_for_location, F.location)
async def process_location(message: Message, state: FSMContext):
    user_lat = message.location.latitude
    user_lon = message.location.longitude
    
    data = await state.get_data()
    jshr = data.get("jshr")
    teacher_name = data.get("teacher_name")
    
    distance = calculate_distance(COLLEGE_LAT, COLLEGE_LON, user_lat, user_lon)
    
    if distance > MAX_DISTANCE_METERS:
        await message.answer(f"❌ Xatolik! Texnikum hududidan uzoqdasiz ({int(distance)} metr).")
        await state.clear()
        return
    
    await state.set_state(AttendanceState.waiting_for_photo)
    await message.answer("Lokatsiya tasdiqlandi ✅ Endi davomat uchun o'z **rasmingizni** yuboring:", reply_markup=None)

@dp.message(AttendanceState.waiting_for_photo, F.photo)
async def process_photo(message: Message, state: FSMContext):
    photo_file_id = message.photo[-1].file_id
    data = await state.get_data()
    jshr = data.get("jshr")
    teacher_name = data.get("teacher_name")
    
    now = datetime.now()
    date_str = now.strftime("%Y-%m-%d")
    time_str = now.strftime("%H:%M")
    status = check_arrival_status()
    
    with open(LOG_FILE, mode="a", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([jshr, teacher_name, date_str, time_str, status, "Keldi"])
        
    await state.clear()
    
    caption_text = (
        f"🔔 **Davomat qayd etildi!**\n\n"
        f"👤 F.I.Sh: {teacher_name}\n"
        f"⏰ Vaqt: {time_str}\n"
        f"📊 Holati: **{status}**"
    )
    
    if status == "Kechikib keldi":
        await message.answer(f"⚠️ Diqqat! Siz belgilangan vaqtdan kechikib keldingiz ({time_str}).")
    else:
        await message.answer(f"Rahmat, {teacher_name}! Kelgan vaqtingiz qayd etildi ✅")
        
    # Xodimning o'ziga rasm va xabar yuborish
    await bot.send_photo(chat_id=message.from_user.id, photo=photo_file_id, caption=caption_text, parse_mode="Markdown")
    
    # 4 ta adminga rasm va xabar yuborish
    for admin_id in ADMIN_IDS:
        try:
            await bot.send_photo(chat_id=admin_id, photo=photo_file_id, caption=caption_text, parse_mode="Markdown")
        except Exception as e:
            logging.error(f"Adminga yuborishda xatolik: {e}")

@dp.message(Command("ketish"))
async def cmd_leave(message: Message, state: FSMContext):
    await state.set_state(AttendanceState.waiting_for_leave)
    await message.answer("Erta ketish uchun PNFL (JSHR) raqamingizni kiriting:")

@dp.message(AttendanceState.waiting_for_leave, F.text)
async def process_leave(message: Message, state: FSMContext):
    text = message.text.strip()
    teacher_name = get_teacher_from_excel(text)
    if not teacher_name or teacher_name == "EXCLUDED":
        teacher_name = "Noma'lum xodim"
        
    now = datetime.now()
    time_str = now.strftime("%H:%M")
    limit_leave_time = now.replace(hour=16, minute=0, second=0, microsecond=0)
    
    await state.clear()
    
    if now < limit_leave_time:
        await message.answer("⚠️ Diqqat! Hali ish vaqti tugamadi (16:00 gacha qolishingiz shart).")
        admin_alert = (
            f"🚨 **ERTA KETISH HARAKATI!**\n\n"
            f"👤 Xodim: **{teacher_name}**\n"
            f"⏰ Vaqt: {time_str}\n"
            f"⚠️ Xodim 16:00 gacha bo'lgan vaqtda binoni tark etmoqchi!"
        )
        for admin_id in ADMIN_IDS:
            try:
                await bot.send_message(admin_id, admin_alert, parse_mode="Markdown")
            except Exception as e:
                logging.error(f"Adminga yuborishda xatolik: {e}")
    else:
        await message.answer("Ish vaqti yakunlangan. Xayrli kun!")

@dp.message(Command("hisobot"))
async def cmd_hisobot(message: Message):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Davomat"
    ws.append(["PNFL", "F.I.Sh.", "Sana", "Vaqt", "Holati"])
    if os.path.exists(LOG_FILE):
        with open(LOG_FILE, mode="r", encoding="utf-8") as f:
            reader = csv.reader(f)
            next(reader, None)
            for row in reader:
                ws.append(row[:5])
    wb.save(EXCEL_REPORT)
    file = FSInputFile(EXCEL_REPORT)
    await message.answer_document(document=file, caption="📊 **Xodimlarning davomat hisoboti**", parse_mode="Markdown")

async def main():
    print("Bot ishga tushdi...")
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
